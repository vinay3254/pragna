import uuid
from pathlib import Path
import chromadb
from pypdf import PdfReader
from app.config import Settings
from app.ollama_client import embed


def chunk_text(text: str, chunk_size: int = 700, overlap: int = 100) -> list[str]:
    text = text.strip()
    if not text:
        return []
    chunks = []
    start = 0
    while start < len(text):
        end = min(start + chunk_size, len(text))
        chunks.append(text[start:end])
        if end == len(text):
            break
        start = end - overlap
    return chunks


def extract_text(file_path: Path) -> str:
    suffix = file_path.suffix.lower()
    if suffix == ".pdf":
        reader = PdfReader(str(file_path))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if suffix == ".docx":
        from docx import Document
        doc = Document(str(file_path))
        parts = [p.text for p in doc.paragraphs if p.text]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    if suffix == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(str(file_path), read_only=True, data_only=True)
        parts = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts)
    if suffix == ".pptx":
        from pptx import Presentation
        prs = Presentation(str(file_path))
        parts = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    parts.append(shape.text)
        return "\n".join(parts)
    return file_path.read_text(encoding="utf-8", errors="ignore")


def get_chroma_collection(settings: Settings):
    client = chromadb.PersistentClient(path=settings.chroma_path)
    return client.get_or_create_collection(
        name="documents", metadata={"hnsw:space": "cosine"}
    )


async def ingest_file(
    file_path: Path,
    filename: str,
    document_id: int,
    collection,
    embed_model: str,
    ollama_url: str,
) -> int:
    text = extract_text(file_path)
    chunks = chunk_text(text)
    if not chunks:
        return 0

    ids, embeddings, metadatas = [], [], []
    for chunk in chunks:
        vector = await embed(chunk, embed_model, ollama_url)
        if not vector:
            vector = [0.0] * 768
        ids.append(str(uuid.uuid4()))
        embeddings.append(vector)
        metadatas.append({"document_id": document_id, "filename": filename})

    collection.add(ids=ids, embeddings=embeddings, documents=chunks, metadatas=metadatas)
    return len(chunks)


async def retrieve(
    query: str,
    collection,
    embed_model: str,
    ollama_url: str,
    top_k: int = 5,
    threshold: float = 0.1,
    where: dict | None = None,
) -> list[dict]:
    if collection.count() == 0:
        return []

    try:
        query_embedding = await embed(query, embed_model, ollama_url)
    except Exception:
        query_embedding = []

    if not query_embedding:
        return _leading_chunks(collection, top_k, where)

    query_kwargs: dict = {
        "query_embeddings": [query_embedding],
        "n_results": min(top_k, collection.count()),
    }
    if where:
        query_kwargs["where"] = where

    results = collection.query(**query_kwargs)

    if not results or not results.get("documents") or not results["documents"][0]:
        return _leading_chunks(collection, top_k, where)

    retrieved = []
    documents = results["documents"][0]
    metadatas = results["metadatas"][0]
    distances = results["distances"][0]
    for doc, meta, distance in zip(documents, metadatas, distances):
        similarity = 1 - distance
        if similarity >= threshold:
            retrieved.append(
                {
                    "document_id": meta["document_id"],
                    "filename": meta["filename"],
                    "snippet": doc[:600],
                    "similarity": similarity,
                }
            )

    return retrieved or _leading_chunks(collection, top_k, where)


def _leading_chunks(collection, top_k: int, where: dict | None) -> list[dict]:
    """Opening chunks of the documents the user explicitly attached.

    Used when similarity search finds nothing -- embeddings unavailable (the
    chunks were stored as zero vectors) or a generic query like "take a look
    at this file" -- so attached documents still reach the model. Without an
    explicit document filter there is nothing to fall back to.
    """
    if not where:
        return []
    got = collection.get(where=where, limit=top_k, include=["documents", "metadatas"])
    return [
        {
            "document_id": meta["document_id"],
            "filename": meta["filename"],
            "snippet": doc[:600],
            "similarity": 0.0,
        }
        for doc, meta in zip(got.get("documents") or [], got.get("metadatas") or [])
    ]

