"""Permissions, imported content, editable canvas and real export artifacts."""

import asyncio
import io
import json
import socket
import zipfile
import pytest
from bs4 import BeautifulSoup
from app import design_service, repository
from app.design_imports import extract_reference, design_context
from app.design_capture import public_address
from app.design_canvas import mutate_canvas
from tests.test_design_workspace import seed_screen


def collaborator(client):
    data = client.post(
        "/api/auth/register",
        json={"email": "collaborator@example.com", "password": "testpassword123"},
    ).json()
    return {"Authorization": f"Bearer {data['access_token']}"}


def test_private_brand_default_and_context(client):
    theme = dict(design_service.DEFAULT_THEME, primary="#334455")
    brand = client.post(
        "/api/design/systems",
        json={
            "name": "Studio",
            "theme": theme,
            "guidelines": "Calm layouts",
            "is_default": True,
            "components": [{"name": "Card", "html": "<article>Brand card</article>"}],
        },
    ).json()
    project = client.post(
        "/api/design/projects", json={"kind": "presentation", "device": "web"}
    ).json()["project"]
    assert (
        project["kind"] == "presentation" and project["theme"]["primary"] == "#334455"
    )
    assert project["design_system_id"] == brand["id"]
    assert "Brand card" in design_context(client.app.state.conn, project)
    headers = collaborator(client)
    assert client.get("/api/design/systems", headers=headers).json() == []
    assert (
        client.put(
            f"/api/design/systems/{brand['id']}",
            headers=headers,
            json={"name": "stolen", "theme": theme},
        ).status_code
        == 404
    )
    assert (
        client.post("/api/design/projects", json={"design_system_id": None}).json()[
            "project"
        ]["design_system_id"]
        is None
    )


def test_roles_and_collaboration_revision(client):
    pid, sid, version = seed_screen(client)
    headers = collaborator(client)
    path = f"/api/design/projects/{pid}"
    client.post(
        path + "/members", json={"email": "collaborator@example.com", "role": "viewer"}
    )
    assert (
        client.get(path, headers=headers).json()["project"]["access_role"] == "viewer"
    )
    assert (
        client.patch(path, headers=headers, json={"name": "blocked"}).status_code == 403
    )
    assert (
        client.post(
            path + "/comments", headers=headers, json={"content": "blocked"}
        ).status_code
        == 403
    )
    assert (
        client.post(
            f"/api/design/screens/{sid}/text",
            headers=headers,
            json={"selector": "h1", "text": "blocked", "version_id": version},
        ).status_code
        == 403
    )
    client.post(
        path + "/members",
        json={"email": "collaborator@example.com", "role": "commenter"},
    )
    assert (
        client.post(
            path + "/comments",
            headers=headers,
            json={"content": "Make title clear", "screen_id": sid, "selector": "h1"},
        ).status_code
        == 200
    )
    client.post(
        path + "/members", json={"email": "collaborator@example.com", "role": "editor"}
    )
    revision = client.get(path + "/revision").json()["revision"]
    assert (
        client.post(
            f"/api/design/screens/{sid}/text",
            headers=headers,
            json={"selector": "h1", "text": "Shared edit", "version_id": version},
        ).status_code
        == 200
    )
    assert client.get(path + "/revision").json()["revision"] != revision
    assert client.delete(path, headers=headers).status_code == 403
    assert (
        client.post(
            path + "/members",
            headers=headers,
            json={"email": "test@example.com", "role": "viewer"},
        ).status_code
        == 403
    )
    member = client.get(path + "/workspace").json()["members"][0]["id"]
    client.delete(path + f"/members/{member}")
    assert client.get(path, headers=headers).status_code == 404


def test_share_privacy_and_revocation(client):
    pid, _, _ = seed_screen(client)
    path = f"/api/design/projects/{pid}"
    client.post(
        path + "/reference", json={"name": "Brief", "content": "Confidential context"}
    )
    client.post(path + "/comments", json={"content": "Private comment"})
    repository.add_design_message(
        client.app.state.conn, pid, "user", "Private conversation"
    )
    token = client.post(path + "/share").json()["token"]
    response = client.get(f"/api/design/shared/{token}", headers={"Authorization": ""})
    assert response.status_code == 200 and "Original title" in response.text
    for private in (
        "Confidential context",
        "Private comment",
        "Private conversation",
        "owner_id",
        "share_token",
        "pragna-design",
    ):
        assert private not in response.text
    client.delete(path + "/share")
    assert client.get(f"/api/design/shared/{token}").status_code == 404


def test_canvas_ids_styles_conflicts_and_history(client):
    _, sid, version = seed_screen(
        client,
        '<main><section id="original"><button data-design-action="toggle" data-design-target="#detail">Open</button><div id="detail" hidden>Details</div></section></main>',
    )
    response = client.post(
        f"/api/design/screens/{sid}/canvas",
        json={"selector": "section", "operation": "duplicate", "version_id": version},
    )
    assert response.status_code == 200
    sections = BeautifulSoup(response.json()["body"], "html.parser").find_all("section")
    assert sections[0]["id"] != sections[1]["id"]
    assert sections[1].button["data-design-target"] == "#" + sections[1].div["id"]
    assert (
        client.post(
            f"/api/design/screens/{sid}/canvas",
            json={
                "selector": "section",
                "operation": "style",
                "styles": {"color": "#112233"},
                "version_id": version,
            },
        ).status_code
        == 409
    )
    latest = response.json()["version_id"]
    assert (
        client.post(
            f"/api/design/screens/{sid}/canvas",
            json={
                "selector": "section",
                "operation": "style",
                "styles": {"background-color": "url(https://example.com)"},
                "version_id": latest,
            },
        ).status_code
        == 400
    )
    grouped = mutate_canvas(response.json()["body"], "section", "group")
    assert "data-design-group" in grouped
    assert "data-design-group" not in mutate_canvas(
        grouped, "[data-design-group]", "ungroup"
    )
    assert (
        client.post(
            f"/api/design/screens/{sid}/restore", json={"version_id": version}
        ).json()["version_id"]
        == version
    )
    assert (
        client.post(
            f"/api/design/screens/{sid}/restore", json={"version_id": latest}
        ).json()["version_id"]
        == latest
    )


def test_office_import_and_archive_bounds():
    from docx import Document
    from pptx import Presentation
    from pptx.util import Inches
    from openpyxl import Workbook

    doc = Document()
    doc.add_paragraph("Brand tone: calm")
    out = io.BytesIO()
    doc.save(out)
    assert "Brand tone: calm" in extract_reference("brief.docx", out.getvalue())[1]
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(
        Inches(1), Inches(1), Inches(5), Inches(1)
    ).text = "Quarterly story"
    out = io.BytesIO()
    deck.save(out)
    assert "Quarterly story" in extract_reference("deck.pptx", out.getvalue())[1]
    book = Workbook()
    book.active.append(["Revenue", 420])
    out = io.BytesIO()
    book.save(out)
    assert "Revenue | 420" in extract_reference("numbers.xlsx", out.getvalue())[1]
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("huge.txt", "x" * (41 * 1024 * 1024))
    with pytest.raises(ValueError, match="expanded"):
        extract_reference("source.zip", out.getvalue())


def test_safe_html_and_media_ranges(client):
    pid, _, _ = seed_screen(client)
    path = f"/api/design/projects/{pid}"
    imported = client.post(
        path + "/import-html",
        json={
            "html": '<style>h1{color:red}</style><h1 onclick="alert(1)">Imported</h1><script>alert(1)</script>'
        },
    ).json()
    assert (
        "onclick" not in imported["body"]
        and "<script>" not in imported["body"]
        and "<style>" in imported["body"]
    )
    source = client.post(
        path + "/sources",
        files={
            "file": (
                "clip.mp4",
                b"\x00\x00\x00\x18ftypmp42" + b"0123456789",
                "video/mp4",
            )
        },
    ).json()
    assert source["media_type"] == "video/mp4"
    asset = f"/api/design/assets/{source['id']}/{source['metadata']['asset_token']}"
    assert client.get(asset + "wrong").status_code == 404
    chunk = client.get(asset, headers={"Range": "bytes=4-7"})
    assert chunk.status_code == 206 and chunk.content == b"ftyp"
    assert client.get(asset, headers={"Range": "bytes=999-"}).status_code == 416


def test_capture_blocks_private_and_mixed_dns(monkeypatch):
    def answers(host, port, *args):
        return [
            (socket.AF_INET, socket.SOCK_STREAM, 6, "", (ip, port))
            for ip in (
                ["8.8.8.8", "127.0.0.1"] if host == "mixed.test" else ["127.0.0.1"]
            )
        ]

    monkeypatch.setattr(socket, "getaddrinfo", answers)
    for url in (
        "http://127.0.0.1",
        "http://mixed.test",
        "file:///etc/passwd",
        "https://user:pass@example.com",
        "http://example.com:8000",
    ):
        with pytest.raises(ValueError):
            asyncio.run(public_address(url))


def test_real_pdf_pptx_and_handoff(client):
    from pathlib import Path
    from playwright.sync_api import sync_playwright

    with sync_playwright() as browser:
        if not Path(browser.chromium.executable_path).is_file():
            pytest.skip("Install Playwright Chromium to verify rendered exports")
    from pptx import Presentation
    from pypdf import PdfReader

    pid, _, _ = seed_screen(
        client,
        '<main class="bg-background p-8"><h1 class="text-3xl font-bold">Editable title</h1><p>Useful content</p></main>',
    )
    from app.design_exports import capture_document
    from PIL import Image

    document = design_service.render_document(
        '<main class="min-h-screen bg-background">Hero</main><footer style="height:100px;background:#ff0000">Footer</footer>',
        design_service.DEFAULT_THEME,
        0,
        False,
    )
    capture = asyncio.run(capture_document(document, 390, 200))
    image = Image.open(io.BytesIO(capture["png"])).convert("RGB")
    assert image.getpixel((195, image.height - 5)) == (255, 0, 0)
    for format in ("pdf", "pptx", "handoff"):
        response = client.get(f"/api/design/projects/{pid}/download/{format}")
        assert response.status_code == 200
        if format == "pdf":
            assert len(PdfReader(io.BytesIO(response.content)).pages) == 1
        elif format == "pptx":
            deck = Presentation(io.BytesIO(response.content))
            assert len(deck.slides) == 1
            assert "Editable title" in "\n".join(
                shape.text for shape in deck.slides[0].shapes if shape.has_text_frame
            )
        else:
            with zipfile.ZipFile(io.BytesIO(response.content)) as bundle:
                assert (
                    "HANDOFF.md" in bundle.namelist()
                    and "vendor/tailwindcss-play-3.4.17.js" in bundle.namelist()
                )
                assert "Editable title" in bundle.read("artboards/01.html").decode()


def test_slides_and_variants_preserve_artboards(client, monkeypatch):
    calls = []

    async def llm(settings, messages, **kwargs):
        calls.append(messages)
        if kwargs.get("stage") == "planning":
            return json.dumps(
                {
                    "screens": [
                        {"name": "Introduction", "purpose": "Context"},
                        {"name": "Next steps", "purpose": "Concrete plan"},
                    ]
                }
            )
        return '<main class="bg-background"><h1>Complete slide</h1></main>'

    monkeypatch.setattr(design_service, "_llm", llm)
    for kind, options in [("presentation", {}), ("prototype", {"variants": 2})]:
        pid = client.post(
            "/api/design/projects", json={"device": "web", "kind": kind}
        ).json()["project"]["id"]
        assert (
            client.post(
                f"/api/design/projects/{pid}/generate",
                json={"prompt": "Present our strategy", **options},
            ).status_code
            == 200
        )
        screens = client.get(f"/api/design/projects/{pid}").json()["screens"]
        assert len(screens) == 2
        if kind == "presentation":
            assert screens[0]["layout"] == {"width": 1280, "height": 720}
    assert any("Build ONLY slide" in str(messages) for messages in calls)


def test_ai_demo_is_bounded_and_private_sources_not_forwarded(client, monkeypatch):
    pid, sid, _ = seed_screen(client)
    client.post(
        f"/api/design/projects/{pid}/reference",
        json={"name": "secret", "content": "Private source text"},
    )
    calls = []

    async def llm(settings, messages, **kwargs):
        calls.append((messages, kwargs))
        return "Prototype response"

    monkeypatch.setattr(design_service, "_llm", llm)
    assert (
        client.post(
            f"/api/design/screens/{sid}/assistant", json={"prompt": "Explain demo"}
        ).json()["content"]
        == "Prototype response"
    )
    assert calls[0][1]["max_tokens"] == 1024 and "Private source text" not in str(calls)


def test_group_multiple_layers_keeps_content_and_anchors(client):
    pid, sid, version = seed_screen(client)
    response = client.post(
        f"/api/design/screens/{sid}/canvas",
        json={
            "selector": "h1",
            "selectors": ["h1", "p"],
            "operation": "group",
            "version_id": version,
        },
    )
    assert response.status_code == 200
    group = BeautifulSoup(response.json()["body"], "html.parser").select_one(
        "[data-design-group]"
    )
    assert group.h1.text == "Original title" and group.p.text == "Some text"
    client.post(
        f"/api/design/projects/{pid}/comments",
        json={"screen_id": sid, "selector": "h1", "content": "Clarify"},
    )
    comment = client.get(f"/api/design/projects/{pid}/workspace").json()["comments"][0]
    assert comment["selector"].startswith("[data-design-anchor=")
    latest = client.get(f"/api/design/projects/{pid}").json()["screens"][0]
    assert (
        BeautifulSoup(latest["body"], "html.parser")
        .select_one(comment["selector"])
        .text
        == "Original title"
    )
    duplicate = client.post(f"/api/design/projects/{pid}/duplicate").json()
    assert (
        duplicate["project"]["id"] != pid
        and "Original title" in duplicate["screens"][0]["body"]
    )


def test_ai_and_voice_controls_are_checked_as_real_interactions():
    body = '<form data-design-ai="#answer"><label>Prompt<input id="prompt" name="prompt" required></label><button type="submit">Ask</button><button type="button" data-design-dictate="#prompt">Speak</button></form><p id="answer" hidden></p><button data-design-speak="#answer">Read answer</button>'
    assert (
        design_service.screen_quality(body, design_service.DEFAULT_THEME, "web")[
            "issues"
        ]
        == []
    )
    broken = design_service.screen_quality(
        body.replace('data-design-dictate="#prompt"', 'data-design-dictate="#missing"'),
        design_service.DEFAULT_THEME,
        "web",
    )
    assert any(issue["code"] == "missing_target" for issue in broken["issues"])
